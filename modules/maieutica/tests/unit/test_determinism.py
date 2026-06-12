"""Unit tests for maieutica.output.determinism — T015.

TDD: tests written BEFORE implementation (RED→GREEN).
"""

from __future__ import annotations

import datetime
import io
import re
import zipfile
from pathlib import Path

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_minimal_xlsx(path: Path) -> None:
    """Create a minimal XLSX file via openpyxl (contains dcterms stamps)."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.save(str(path))


# ---------------------------------------------------------------------------
# finalize_xlsx
# ---------------------------------------------------------------------------


class TestFinalizeXlsx:
    def test_dcterms_modified_is_pinned(self, tmp_path: Path) -> None:
        """After finalize_xlsx, docProps/core.xml dcterms:modified equals the pinned value."""
        from maieutica.output.determinism import finalize_xlsx

        xlsx = tmp_path / "test.xlsx"
        _make_minimal_xlsx(xlsx)

        when = datetime.datetime(2026, 1, 1, 0, 0, 0)
        finalize_xlsx(xlsx, when)

        with zipfile.ZipFile(xlsx, "r") as zf:
            core_xml = zf.read("docProps/core.xml").decode("utf-8")

        modified_re = re.compile(r"<dcterms:modified[^>]*>([^<]+)</dcterms:modified>")
        match = modified_re.search(core_xml)
        assert match, "dcterms:modified not found in core.xml"
        assert match.group(1) == "2026-01-01T00:00:00Z"

    def test_dcterms_created_is_pinned(self, tmp_path: Path) -> None:
        """After finalize_xlsx, docProps/core.xml dcterms:created equals the pinned value."""
        from maieutica.output.determinism import finalize_xlsx

        xlsx = tmp_path / "test_created.xlsx"
        _make_minimal_xlsx(xlsx)

        when = datetime.datetime(2026, 1, 1, 0, 0, 0)
        finalize_xlsx(xlsx, when)

        with zipfile.ZipFile(xlsx, "r") as zf:
            core_xml = zf.read("docProps/core.xml").decode("utf-8")

        created_re = re.compile(r"<dcterms:created[^>]*>([^<]+)</dcterms:created>")
        match = created_re.search(core_xml)
        assert match, "dcterms:created not found in core.xml"
        assert match.group(1) == "2026-01-01T00:00:00Z"

    def test_two_calls_produce_byte_identical_output(self, tmp_path: Path) -> None:
        """Two sequential finalize_xlsx calls on separate copies produce identical bytes."""
        from maieutica.output.determinism import finalize_xlsx

        xlsx_a = tmp_path / "a.xlsx"
        xlsx_b = tmp_path / "b.xlsx"
        _make_minimal_xlsx(xlsx_a)
        _make_minimal_xlsx(xlsx_b)

        when = datetime.datetime(2026, 3, 15, 12, 0, 0)
        finalize_xlsx(xlsx_a, when)
        finalize_xlsx(xlsx_b, when)

        assert xlsx_a.read_bytes() == xlsx_b.read_bytes(), (
            "finalize_xlsx must produce byte-identical output for identical inputs"
        )

    def test_byte_identical_when_created_timestamps_differ(self, tmp_path: Path) -> None:
        """Two xlsx with DIFFERENT openpyxl created timestamps finalize byte-identical.

        Simulates the real flake: build #1 and build #2 land in different wall-clock
        seconds, so openpyxl writes different <dcterms:created>.
        finalize_xlsx must normalise BOTH so the bytes match.
        """
        from maieutica.output.determinism import finalize_xlsx

        xlsx_a = tmp_path / "created_a.xlsx"
        xlsx_b = tmp_path / "created_b.xlsx"
        _make_minimal_xlsx(xlsx_a)
        _make_minimal_xlsx(xlsx_b)

        def _inject_core(path: Path, ts: str) -> None:
            with zipfile.ZipFile(path, "r") as src:
                members = [
                    (i.filename, src.read(i.filename), i.compress_type)
                    for i in src.infolist()
                ]
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
                for name, data, ct in members:
                    if name == "docProps/core.xml":
                        text = data.decode("utf-8")
                        text = re.sub(
                            r"(<dcterms:created[^>]*>)[^<]+(</dcterms:created>)",
                            rf"\g<1>{ts}\g<2>",
                            text,
                        )
                        text = re.sub(
                            r"(<dcterms:modified[^>]*>)[^<]+(</dcterms:modified>)",
                            rf"\g<1>{ts}\g<2>",
                            text,
                        )
                        data = text.encode("utf-8")
                    dst.writestr(name, data, compress_type=ct)
            path.write_bytes(buf.getvalue())

        _inject_core(xlsx_a, "2026-05-31T08:00:01Z")
        _inject_core(xlsx_b, "2026-05-31T08:00:02Z")

        when = datetime.datetime(2026, 1, 1, 0, 0, 0)
        finalize_xlsx(xlsx_a, when)
        finalize_xlsx(xlsx_b, when)

        assert xlsx_a.read_bytes() == xlsx_b.read_bytes()

    def test_zip_entries_use_fixed_date_time(self, tmp_path: Path) -> None:
        """All zip entries use the fixed date_time (1980, 1, 1, 0, 0, 0)."""
        from maieutica.output.determinism import finalize_xlsx

        xlsx = tmp_path / "fixed_date.xlsx"
        _make_minimal_xlsx(xlsx)
        finalize_xlsx(xlsx, datetime.datetime(2026, 1, 1))

        with zipfile.ZipFile(xlsx, "r") as zf:
            for info in zf.infolist():
                assert info.date_time == (1980, 1, 1, 0, 0, 0), (
                    f"Entry {info.filename!r} has date_time {info.date_time}, "
                    "expected (1980,1,1,0,0,0)"
                )


# ---------------------------------------------------------------------------
# dump_yaml
# ---------------------------------------------------------------------------


class TestDumpYaml:
    def test_byte_identical_across_two_calls(self) -> None:
        """dump_yaml is deterministic: two calls with same input produce identical output."""
        from maieutica.output.determinism import dump_yaml

        obj = {"z_key": "value", "a_key": 42, "m_key": [1, 2, 3]}
        r1 = dump_yaml(obj)
        r2 = dump_yaml(obj)
        assert r1 == r2

    def test_unicode_preserved(self) -> None:
        """Korean characters are preserved, not escaped."""
        from maieutica.output.determinism import dump_yaml

        obj = {"제목": "기말고사", "항목": ["호흡계통", "근육계통"]}
        result = dump_yaml(obj)
        assert "기말고사" in result
        assert "호흡계통" in result
        assert "\\u" not in result

    def test_sorted_keys(self) -> None:
        """Keys are sorted alphabetically for determinism."""
        from maieutica.output.determinism import dump_yaml

        obj = {"z": 1, "a": 2, "m": 3}
        result = dump_yaml(obj)
        lines = [line for line in result.splitlines() if ":" in line]
        keys = [line.split(":")[0].strip() for line in lines]
        assert keys == sorted(keys), f"Keys not sorted: {keys}"

    def test_roundtrip(self) -> None:
        """dump_yaml output can be parsed back to the original object."""
        import yaml as pyyaml
        from maieutica.output.determinism import dump_yaml

        obj = {"key": "value", "nested": {"a": 1, "b": 2}, "list": [1, 2, 3]}
        dumped = dump_yaml(obj)
        restored = pyyaml.safe_load(dumped)
        assert restored == obj

    def test_ends_with_newline(self) -> None:
        """YAML output ends with exactly one newline."""
        from maieutica.output.determinism import dump_yaml

        result = dump_yaml({"x": 1})
        assert result.endswith("\n"), "YAML must end with a newline"
        assert not result.endswith("\n\n"), "YAML must not end with double newline"


# ---------------------------------------------------------------------------
# xls_byte_gate (NEW — maieutica-specific, R1)
# ---------------------------------------------------------------------------


class TestXlsByteGate:
    """Tests for the .xls byte-determinism gate helper."""

    def test_gate_passes_for_deterministic_writer(self, tmp_path: Path) -> None:
        """gate_xls_deterministic passes when the writer is byte-identical on 2 runs."""
        from maieutica.output.determinism import gate_xls_deterministic

        def deterministic_writer(path: Path) -> None:
            # Fixed bytes, no timestamps — always identical
            path.write_bytes(b"XLS\x00static content\x00")

        # Should not raise
        gate_xls_deterministic(deterministic_writer, work_dir=tmp_path)

    def test_gate_fails_for_nondeterministic_writer(self, tmp_path: Path) -> None:
        """gate_xls_deterministic raises AssertionError when writer embeds a timestamp."""

        import pytest
        from maieutica.output.determinism import gate_xls_deterministic

        call_count = [0]

        def nondeterministic_writer(path: Path) -> None:
            # Embed a counter (simulates a timestamp changing between calls)
            call_count[0] += 1
            path.write_bytes(f"XLS content call={call_count[0]}".encode())

        with pytest.raises(AssertionError):
            gate_xls_deterministic(nondeterministic_writer, work_dir=tmp_path)

    def test_gate_cleans_up_temp_files(self, tmp_path: Path) -> None:
        """gate_xls_deterministic leaves no temp files after completion."""
        from maieutica.output.determinism import gate_xls_deterministic

        def writer(path: Path) -> None:
            path.write_bytes(b"fixed")

        gate_xls_deterministic(writer, work_dir=tmp_path)
        # No .xls temp files should remain
        remaining = list(tmp_path.glob("*.xls"))
        assert remaining == [], f"Leftover temp files: {remaining}"

    def test_gate_cleans_up_on_failure(self, tmp_path: Path) -> None:
        """gate_xls_deterministic leaves no temp files even when assertion fails."""
        import pytest
        from maieutica.output.determinism import gate_xls_deterministic

        counter = [0]

        def nd_writer(path: Path) -> None:
            counter[0] += 1
            path.write_bytes(str(counter[0]).encode())

        with pytest.raises(AssertionError):
            gate_xls_deterministic(nd_writer, work_dir=tmp_path)

        remaining = list(tmp_path.glob("*.xls"))
        assert remaining == [], f"Leftover temp files after failure: {remaining}"


# ---------------------------------------------------------------------------
# T013 smoke: templates load and contain expected placeholders
# ---------------------------------------------------------------------------


class TestConfigTemplates:
    def test_quiz_column_map_loads_and_has_11_columns(self) -> None:
        """quiz_column_map.yaml loads and contains the 11 LMS column keys."""
        from pathlib import Path

        import yaml

        template_path = (
            Path(__file__).parent.parent.parent
            / "templates"
            / "quiz_column_map.yaml"
        )
        assert template_path.exists(), f"Template not found: {template_path}"
        data = yaml.safe_load(template_path.read_text(encoding="utf-8"))

        columns = data.get("columns", {})
        headers = {v["header"] for v in columns.values()}
        expected_headers = {
            "문제번호",
            "문제내용",
            "예상주차",
            "보기1",
            "보기2",
            "보기3",
            "보기4",
            "보기5",
            "답안",
            "답안설명",
            "문항유형",
        }
        assert headers == expected_headers, (
            f"Mismatch — missing: {expected_headers - headers}, "
            f"extra: {headers - expected_headers}"
        )

    def test_quiz_column_map_cell_type_contract(self) -> None:
        """quiz_column_map.yaml encodes LMS cell-type contract (SC-003)."""
        from pathlib import Path

        import yaml

        template_path = (
            Path(__file__).parent.parent.parent
            / "templates"
            / "quiz_column_map.yaml"
        )
        data = yaml.safe_load(template_path.read_text(encoding="utf-8"))
        columns = data["columns"]

        # 문제번호 must be "number"
        assert columns["item_no"]["cell_type"] == "number"
        # 답안 must be "text" (SC-003 trap: "3" not 3)
        assert columns["answer"]["cell_type"] == "text"
        # 예상주차 must be "text" with zero_pad3
        assert columns["week"]["cell_type"] == "text"
        assert columns["week"].get("format") == "zero_pad3"
        # 문항유형 constant must be "002"
        assert columns["question_type_code"]["constant"] == "002"

    def test_prompt_quiz_contains_required_placeholders(self) -> None:
        """prompt_quiz.txt contains all documented {named} placeholders."""
        from pathlib import Path

        template_path = (
            Path(__file__).parent.parent.parent
            / "templates"
            / "prompt_quiz.txt"
        )
        assert template_path.exists()
        content = template_path.read_text(encoding="utf-8")

        required = {
            "{chapter}",
            "{chapter_no}",
            "{section}",
            "{week}",
            "{quiz_count}",
            "{textbook_context}",
            "{key_concept}",
            "{slot_id}",
            "{question_type}",
        }
        for placeholder in required:
            assert placeholder in content, (
                f"Missing placeholder {placeholder} in prompt_quiz.txt"
            )

    def test_prompt_formative_contains_required_placeholders(self) -> None:
        """prompt_formative.txt contains all documented {named} placeholders."""
        from pathlib import Path

        template_path = (
            Path(__file__).parent.parent.parent
            / "templates"
            / "prompt_formative.txt"
        )
        assert template_path.exists()
        content = template_path.read_text(encoding="utf-8")

        required = {
            "{chapter}",
            "{chapter_no}",
            "{section}",
            "{week}",
            "{formative_count}",
            "{textbook_context}",
            "{key_concept}",
            "{slot_id}",
            "{topic}",
        }
        for placeholder in required:
            assert placeholder in content, (
                f"Missing placeholder {placeholder} in prompt_formative.txt"
            )
